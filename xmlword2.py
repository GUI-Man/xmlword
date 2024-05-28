import os
from docx import Document
import xlwt
import xlrd
import queue
class Mulu:
    def __init__(self,name,text):

        self.name=name
        self.depth=-1
        self.weightMax=0
        self.weightMin=0
        self.text=text
        self.abovename=[]
        self.belowMulu=[]
        self.filename=""
        self.xmlfile=""
        self.FatherMulu=None
class Main:
    def __init__(self,filename,xmlfile):
        self.filename=filename
        self.xmlfile=xmlfile
        self.index=0
        self.indexy=0
    def AddSon(self,father,son):
        son.FatherMulu=father
        son.abovename=son.FatherMulu.abovename.copy()
        son.depth=father.depth+1
        son.abovename.append(father.name)
        father.belowMulu.append(son)
    def FindFatherByName(self,son,name):
        if(son.name==name):
            return son
        elif(son.FatherMulu==None):
            return False
        else:
            return self.FindFatherByName(son.FatherMulu,name)
    def VisitDetailedName(self,here):
        print("    " * self.index + here.text + "(" + str(here.depth) + ")" + "weightMin(" + str(
            here.weightMin) + ")" + "weightMax(" + str(here.weightMax) + ")")
        self.index += 1
        for i in here.belowMulu:
            self.VisitDetailedName(i)
        self.index -= 1
    def writeInto(self,Heading,table2):
        if(Heading.weightMin==Heading.weightMax):
            if(Heading.depth>=0):
                table2.write(Heading.weightMin,Heading.depth,Heading.text)
        elif(Heading.weightMax>Heading.weightMin):
            if(Heading.depth>=0):
                table2.write_merge(Heading.weightMin,Heading.weightMax,Heading.depth,Heading.depth,Heading.text)
        for i in Heading.belowMulu:
            self.writeInto(i,table2)
    def VisitName(self,here):

        if(here.belowMulu==[]):
            here.weightMax=self.indexy
            here.weightMin=self.indexy
            self.indexy+=1
        print("    " * self.index + here.text)
        self.index+=1
        for i in here.belowMulu:
            self.VisitName(i)
        if(len(here.belowMulu)!=0):
            here.weightMax=here.belowMulu[-1].weightMax
            here.weightMin=here.belowMulu[0].weightMin
        self.index-=1
    def main(self):
        doc=Document(self.xmlfile)
        heading_list=[]
        Head=Mulu("ALL","开始")
        Now=Head
        for paragrhaph in doc.paragraphs:
            paraname=paragrhaph.style.name
            if(paraname!="Normal"):
                if(paraname not in Now.abovename):
                    thisMulu=Mulu(paraname,paragrhaph.text)
                    self.AddSon(Now,thisMulu)
                    Now=thisMulu
                elif(paraname in Now.abovename):
                    Now=self.FindFatherByName(Now.FatherMulu,paraname)
                    if(Now==False):
                        print("?")
                        exit()
                    else:
                        thisMulu=Mulu(paraname,paragrhaph.text)
                        self.AddSon(Now.FatherMulu,thisMulu)
                        Now=thisMulu
        print("pause")
        self.index=0
        self.indexy=0
        self.VisitName(Head)
        f2 = open(self.filename, "w+", encoding='utf-8')
        df2 = xlwt.Workbook(f2)
        table2 = df2.add_sheet('test')
        self.index=0
        self.indexy=0
        self.VisitDetailedName(Head)
        self.writeInto(Head,table2)
        df2.save(self.filename)
        f2.close()
# if __name__ == '__main__':
#
#     A=Main(r".\result2.xls",r"D:\学习\爬虫\XMLWord\doc\M3接口分册20240508.docx")
#     A.main()

